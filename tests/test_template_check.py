import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import openpyxl

from config.config_loader import ConfigError
from core.context import TaskContext
from workflows.jde_flow.property_tax import PropertyTaxFlow, PendingModule
from workflows.jde_flow.template_check import check_sheet


class TemplateCheckTests(unittest.TestCase):
    def sheet(self, count):
        book = openpyxl.Workbook()
        self.addCleanup(book.close)
        sheet = book.active
        sheet.title = "房产税（从租）计算底稿"
        sheet["B7"] = "BU"
        sheet["H7"] = "借项\n金额"
        sheet["P4"] = "分类帐类型 1 金额"
        for i in range(count):
            sheet.cell(9+i, 2).value = 250104
            sheet.cell(9+i, 3).value = 100
            sheet.cell(9+i, 4).value = 12
            sheet.cell(5+i, 15).value = f"250104.{i}"
            sheet.cell(5+i, 16).value = -100
        fixed = 9+count
        sheet.cell(fixed, 2).value = 250108
        sheet.cell(fixed, 3).value = 0
        sheet.cell(fixed, 4).value = 0
        sheet.cell(fixed+2, 2).value = "合计"
        sheet.cell(fixed+2, 3).value = count*100
        sheet.cell(fixed+2, 4).value = count*12
        sheet["G8"] = "250104.917013"
        sheet["H8"] = count*12
        sheet.cell(5+count, 16).value = -count*100
        return book, sheet

    def test_dynamic_rows_and_signed_check_total(self):
        for count in (1, 6, 25):
            with self.subTest(count=count):
                _, sheet = self.sheet(count)
                results = check_sheet(sheet)
                self.assertTrue(all(item["通过"] for item in results))
                self.assertEqual(results[0]["位置"]["税额合计"], f"D{11+count}")
                self.assertEqual(results[1]["位置"]["Check合计"], f"P{5+count}")

    def test_fixed_bu_is_not_deducted_and_mismatch_fails(self):
        _, sheet = self.sheet(3)
        sheet["C12"] = 10
        sheet["C14"] = 310
        result = check_sheet(sheet)
        self.assertFalse(result[1]["通过"])
        self.assertEqual(result[1]["差额"], "10")

    def test_missing_formula_cache_rejected(self):
        _, sheet = self.sheet(2)
        sheet["H8"] = None
        with self.assertRaisesRegex(ConfigError, "无法取数"):
            check_sheet(sheet)

    @patch("workflows.jde_flow.property_tax.recalculate_file")
    def test_failure_saved_and_pass_does_not_authorize_journal(self, recalc):
        book, sheet = self.sheet(2)
        context = TaskContext()
        class Config:
            def get(self, key, default=None):
                return default
        context.config = Config()
        flow = PropertyTaxFlow(context)
        with TemporaryDirectory() as folder:
            path = Path(folder) / "output.xlsx"
            book.save(path)
            with self.assertRaisesRegex(PendingModule, "核对通过"):
                flow.compare_values(path)
            self.assertEqual(context.get("property_tax_check_status"), "passed")
            self.assertIsNone(context.get("journal_required"))
            sheet["H8"] = 23
            book.save(path)
            with self.assertRaisesRegex(ConfigError, "核对失败"):
                flow.compare_values(path)
            self.assertEqual(context.get("property_tax_check_status"), "failed")
            self.assertTrue(context.get("property_tax_check_report").is_file())
            self.assertEqual(recalc.call_count, 2)
