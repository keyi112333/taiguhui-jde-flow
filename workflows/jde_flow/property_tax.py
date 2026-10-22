"""房产税模块二至六；未确认的业务规则明确阻断后续处理。"""
import json
import re
from copy import copy
from pathlib import Path
from tempfile import NamedTemporaryFile

import openpyxl
from openpyxl.workbook.properties import CalcProperties

from config.config_loader import ConfigError
from log.loggingConfig import logger
from workflows.jde_flow.template_update import update_workbook
from workflows.jde_flow.template_check import recalculate_file, check_sheet


HEADER_DEFINITION = Path(__file__).resolve().parents[2] / "config/property_tax_ledger_headers.json"


class PendingModule(Exception):
    """流程等待录制、数据或业务规则，不代表已完成。"""


class PropertyTaxFlow:
    def __init__(self, context):
        self.context = context
        self.config = context.config

    def get(self, key, default=None):
        return self.config.get(f"房产税配置.{key}", default)

    def export_ledger(self):
        """模块二：后续录制导出后，返回下载文件路径。"""
        supplied = self.context.get("jde_ledger_path") or self.get("明细账文件")
        if not supplied:
            raise PendingModule("模块二待录制；可配置 房产税配置.明细账文件 测试模块三")
        path = Path(supplied)
        if not path.is_absolute():
            path = self.config.get_path("input") / path
        if not path.is_file():
            raise ConfigError(f"明细账文件不存在：{path}")
        return path

    def create_template(self, ledger_path):
        """模块三：保留模板其他页及公式，按标题写入明细。"""
        input_dir = self.config.get_path("input")
        template_name = self.get("模板文件")
        if template_name:
            template = Path(template_name)
            if not template.is_absolute():
                template = input_dir / template
        else:
            candidates = [p for p in input_dir.glob("*房产税计税表*.xlsx")
                          if not p.name.startswith("~$") and p.resolve() != Path(ledger_path).resolve()]
            if len(candidates) != 1:
                raise ConfigError("输入目录需有唯一包含‘房产税计税表’的 xlsx 模板，或配置 模板文件")
            template = candidates[0]
        company, period = self.get("公司"), self.get("会计期间")
        for label, value in (("公司", company), ("会计期间", period)):
            if value is None or not str(value).strip() or re.search(r'[<>:"/\\|?*\x00-\x1f]', str(value)):
                raise ConfigError(f"房产税配置.{label} 必填且不能包含文件名非法字符")
        output_dir = Path(self.get("输出目录") or self.config.get_path("output"))
        if not output_dir.is_absolute():
            output_dir = self.config.get_path("output") / output_dir
        destination = output_dir / f"房产税计税表 {str(company).strip()} {str(period).strip()}.xlsx"
        if destination.exists() or destination.resolve() in (template.resolve(), Path(ledger_path).resolve()):
            raise ConfigError(f"输出文件已存在或与输入重合：{destination}")
        mapping = self.get("标题映射", "")
        try:
            mapping = json.loads(mapping) if mapping else {}
        except (TypeError, ValueError) as exc:
            raise ConfigError("标题映射须为 JSON 对象：源标题 -> 模板标题") from exc
        if not isinstance(mapping, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in mapping.items()):
            raise ConfigError("标题映射须为字符串到字符串的 JSON 对象")
        definition = json.loads(HEADER_DEFINITION.read_text(encoding="utf-8"))
        expected = [column["template_header"] for column in definition["columns"]]
        source = openpyxl.load_workbook(ledger_path, data_only=True)
        target = None
        temp_path = None
        try:
            target = openpyxl.load_workbook(template)
            sheet_name = self.get("明细账工作表") or source.sheetnames[0]
            source_sheet = source[sheet_name]
            target_sheet = target[definition["sheet"]]
            source_row = int(self.get("明细账标题行", 1))
            target_row = int(self.get("模板标题行", definition["header_row"]))
            if min(source_row, target_row) < 1:
                raise ConfigError("标题行必须为正整数")
            def headers(sheet, row):
                result = {}
                for cell in sheet[row]:
                    if cell.value is None or not str(cell.value).strip():
                        continue
                    title = str(cell.value).strip()
                    if title in result:
                        raise ConfigError(f"{sheet.title} 标题重复：{title}")
                    result[title] = cell.column
                return result
            src = headers(source_sheet, source_row)
            dst = headers(target_sheet, target_row)
            if set(dst) != set(expected):
                raise ConfigError(f"模板标题与记录不符；缺少：{sorted(set(expected) - set(dst))}；"
                                  f"多出：{sorted(set(dst) - set(expected))}")
            # JSON 中确认的源标题为默认别名；运行配置优先，同名标题自动匹配。
            aliases = {column["source_header"]: column["template_header"]
                       for column in definition["columns"] if column["source_header"]}
            aliases.update(mapping)
            pairs = {title: aliases.get(title, title) for title in src
                     if title in aliases or title in dst}
            if not pairs:
                raise ConfigError("没有匹配的标题，请补充标题映射")
            if len(set(pairs.values())) != len(pairs):
                raise ConfigError("多个源标题不能写入同一个模板标题")
            for origin, name in pairs.items():
                if origin not in src or name not in dst:
                    raise ConfigError(f"标题不存在：{origin} -> {name}")
            # 不默默丢弃源列；标题尚未明确时应先补齐映射。
            if set(src) - set(pairs):
                raise ConfigError(f"源标题尚未映射：{sorted(set(src) - set(pairs))}")
            if set(dst) - set(pairs.values()):
                raise ConfigError(f"明细账缺少模板所需列：{sorted(set(dst) - set(pairs.values()))}")
            count = 0
            # 保存首条明细格式，扩展新行时延续日期、金额等显示格式。
            styles = {col: copy(target_sheet.cell(target_row + 1, col)._style)
                      for col in dst.values()}
            for col in (dst[name] for name in pairs.values()):
                for row in range(target_row + 1, target_sheet.max_row + 1):
                    target_sheet.cell(row, col).value = None
            for row in source_sheet.iter_rows(min_row=source_row + 1, values_only=True):
                if all(value is None or (isinstance(value, str) and not value.strip()) for value in row):
                    continue
                count += 1
                for origin, name in pairs.items():
                    value = row[src[origin] - 1]
                    cell = target_sheet.cell(target_row + count, dst[name])
                    cell._style = copy(styles[dst[name]])
                    cell.value = value
                    if isinstance(value, str):
                        cell.data_type = "s"
            if not count:
                raise ConfigError("明细账没有数据，未生成输出")
            target.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)
            output_dir.mkdir(parents=True, exist_ok=True)
            with NamedTemporaryFile(dir=output_dir, suffix=".xlsx", delete=False) as tmp:
                temp_path = Path(tmp.name)
            target.save(temp_path)
            temp_path.rename(destination)
            self.context.set("property_tax_path", destination)
            self.context.set("jde_ledger_rows", count)
            self.context.set("jde_ledger_mapping", pairs)
            return destination
        finally:
            source.close()
            if target:
                target.close()
            if temp_path and temp_path.exists():
                temp_path.unlink()

    def update_template(self, path):
        """模块四：更新 BU 汇总；公式由 Excel 打开时重算。"""
        mode = self.get("收入科目匹配方式", "精确")
        if mode not in ("精确", "包含子科目"):
            raise ConfigError("收入科目匹配方式须为 精确 或 包含子科目")
        book = openpyxl.load_workbook(path)
        temp_path = None
        try:
            records = update_workbook(book, include_children=mode == "包含子科目")
            book.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)
            with NamedTemporaryFile(dir=Path(path).parent, suffix=".xlsx", delete=False) as tmp:
                temp_path = Path(tmp.name)
            book.save(temp_path)
            temp_path.replace(path)
            self.context.set("property_tax_bu_totals", records)
        finally:
            book.close()
            if temp_path and temp_path.exists():
                temp_path.unlink()

    def compare_values(self, path):
        """模块五：重算并核对；是否写分录仍待业务规则确认。"""
        self.context.set("property_tax_check_status", "running")
        try:
            recalculate_file(path, self.get("重算引擎", "Excel.Application"))
            book = openpyxl.load_workbook(path, data_only=True)
            try:
                results = check_sheet(book["房产税（从租）计算底稿"],
                                      tolerance=self.get("核对容差", "0.00"))
            finally:
                book.close()
            passed = all(item["通过"] for item in results)
            self.context.set("property_tax_check_results", results)
            self.context.set("property_tax_check_status", "passed" if passed else "failed")
            report = Path(path).with_suffix(".核对结果.json")
            report.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
            self.context.set("property_tax_check_report", report)
            if not passed:
                raise ConfigError("模板核对失败，详见核对结果：" + str(report))
        except Exception:
            self.context.set("property_tax_check_status", "failed")
            raise
        raise PendingModule("模板核对通过；是否进入 JDE 写入分录的规则待补充")

    def write_journal(self, path):
        """模块六：待补充分录字段及 JDE 操作。"""
        raise PendingModule("模块六待补充：JDE 分录写入操作")

    def run(self):
        self.context.set("property_tax_status", "running")
        try:
            path = self.create_template(self.export_ledger())
            self.update_template(path)
            decision = self.compare_values(path)
            if type(decision) is not bool:
                raise ConfigError("模块五必须明确返回是否需要分录")
            self.context.set("journal_required", decision)
            if decision:
                self.write_journal(path)
            self.context.set("property_tax_status", "completed")
        except PendingModule as exc:
            self.context.set("property_tax_status", "pending")
            self.context.set("property_tax_pending_reason", str(exc))
            logger.warning(str(exc))
        except Exception:
            self.context.set("property_tax_status", "failed")
            raise
